from openpyxl import load_workbook
from django.core.management.base import BaseCommand
from django.db import transaction
from rooms.models import Room, RoomSpace


class Command(BaseCommand):
    help = 'Import rooms from MUSHIA.xlsx'

    def add_arguments(self, parser):
        parser.add_argument('excel_path', type=str, help='Path to MUSHIA.xlsx')

    @transaction.atomic
    def handle(self, *args, **kwargs):
        wb = load_workbook(kwargs['excel_path'], data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        if len(rows) < 3:
            return

        headers = rows[0]
        floor_col_indices = [1, 4, 7, 10, 13, 16]
        existing_room_nums = set(Room.objects.values_list('room_number', flat=True))
        rooms_to_create = []
        seen = set()

        for col in floor_col_indices:
            if col >= len(headers) or not headers[col]:
                continue
            floor_name = str(headers[col]).strip()

            for row in rows[2:]:
                if col + 1 >= len(row):
                    continue
                room_num, room_type = row[col], row[col + 1]
                if room_num is None or room_type is None:
                    continue

                room_num_str = str(room_num).strip()
                room_type_str = str(room_type).strip()
                if not room_num_str or room_num_str in existing_room_nums or room_num_str in seen:
                    continue
                seen.add(room_num_str)

                upper_type = room_type_str.upper()
                amenities = []
                if 'AC' in upper_type:
                    amenities.append('AC')
                if 'LARGE' in upper_type or 'BIG' in upper_type:
                    amenities.append('Big')

                if 'RESERVED' in upper_type:
                    amenities.append('Reserved')
                    capacity = 1
                else:
                    capacity = int(room_type_str[0]) if room_type_str[0].isdigit() else 1

                if capacity >= 3 and 'Big' not in amenities:
                    amenities.append('Big')

                rooms_to_create.append(
                    Room(
                        room_number=room_num_str,
                        floor=floor_name,
                        room_type=room_type_str,
                        capacity=capacity,
                        price_per_space=0.00,
                        amenities=amenities,
                    )
                )

        rooms_created = 0
        spaces_created = 0
        if rooms_to_create:
            created_rooms = Room.objects.bulk_create(rooms_to_create)
            letters = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
            spaces_to_create = [
                RoomSpace(room=room, space_identifier=letters[j])
                for room in created_rooms
                for j in range(room.capacity)
            ]
            RoomSpace.objects.bulk_create(spaces_to_create)
            rooms_created = len(created_rooms)
            spaces_created = len(spaces_to_create)

        self.stdout.write(self.style.SUCCESS(f"Successfully created {rooms_created} rooms and {spaces_created} spaces!"))
