import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_documentation():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Styles
    COLOR_PRIMARY = RGBColor(42, 40, 39)       # #2A2827 (Charcoal)
    COLOR_ACCENT = RGBColor(190, 150, 0)       # Deep Gold / Yellow accent
    COLOR_MUTED = RGBColor(100, 100, 100)      # Gray
    COLOR_TEXT = RGBColor(50, 50, 50)          # Body Dark
    COLOR_CODE = RGBColor(20, 80, 140)         # Code Blue

    # ------------------ TITLE & COVER HEADER ------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(20)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("MUSHIA HOSTEL RESERVATION PLATFORM")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run("Frontend Developer Integration Guide & System Architecture Specification")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = COLOR_ACCENT
    run_sub.font.bold = True

    # Metadata Box
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(24)
    run_meta = p_meta.add_run("Version 1.0 (Production-Ready)  |  KNUST Kumasi  |  October 2026\nTech Stack: React 19 + Vite 8 + Django REST Framework + Paystack")
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(9.5)
    run_meta.font.italic = True
    run_meta.font.color.rgb = COLOR_MUTED

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Helper function for Section Headings
    def add_section_header(title, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title)
        run.font.name = "Arial"
        run.font.bold = True
        if level == 1:
            run.font.size = Pt(15)
            run.font.color.rgb = COLOR_PRIMARY
        elif level == 2:
            run.font.size = Pt(12)
            run.font.color.rgb = COLOR_ACCENT
        else:
            run.font.size = Pt(10.5)
            run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_body_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Arial"
            r_bold.font.size = Pt(10)
            r_bold.font.bold = True
            r_bold.font.color.rgb = COLOR_PRIMARY
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_TEXT
        return p

    def add_bullet_p(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Arial"
            r_bold.font.size = Pt(10)
            r_bold.font.bold = True
            r_bold.font.color.rgb = COLOR_PRIMARY
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_TEXT
        return p

    def add_code_block(code_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F7F7F8")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(code_text)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = COLOR_CODE
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ------------------ SECTION 1: SYSTEM OVERVIEW ------------------
    add_section_header("1. Executive Project Overview & Architecture")
    add_body_p(
        "This platform is an enterprise-grade hostel reservation system engineered for Mushia Hostel (KNUST, Kumasi). "
        "It provides real-time room and bed-space discovery, atomic gender policy enforcement, concurrency-safe "
        "15-minute temporary holds, automated Paystack payment verification, and privacy-protected roommate contact visibility."
    )
    add_body_p(
        "The codebase was built with strict separation of concerns, ensuring that the original cloned Vite React UI styling, "
        "palette (#2A2827 charcoal, #FEFB58 yellow, #F4EFE7 light cream), typography, and animations remained 100% pristine. "
        "All backend interactions connect through a single centralized state layer (HostelContext.tsx and api.ts)."
    )

    add_bullet_p("React 19, Vite 8, TypeScript, Tailwind CSS v4, Lucide Icons, Axios.", "Frontend Client: ")
    add_bullet_p("Django 6.1, Django REST Framework (DRF), SimpleJWT with HttpOnly cookies, uv package manager.", "Backend Service: ")
    add_bullet_p("SQLite (Local Dev / Staging) & PostgreSQL (Production-ready schema).", "Database: ")
    add_bullet_p("Paystack Checkout (GHS currency), Webhooks with HMAC-SHA512 verification.", "Payment Gateway: ")
    add_bullet_p("Nginx reverse proxy for single-domain routing, zero CORS issues, and static/media caching.", "Reverse Proxy: ")

    # ------------------ SECTION 2: DATABASE SCHEMA ------------------
    add_section_header("2. Database Schema & Data Models")
    add_body_p(
        "The backend database is structured into 5 relational models mapped via Django ORM. "
        "All 102 physical rooms and 193 individual bed spaces from MUSHIA.xlsx are seeded and indexed."
    )

    # Model 1: Student
    add_section_header("Model: Student (Custom User)", level=2)
    add_body_p("Extends AbstractBaseUser and PermissionsMixin. Authenticates using email as the primary credential.")

    table_student = doc.add_table(rows=1, cols=4)
    table_student.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Field Name", "Type / Nullable", "Constraints / Choices", "Description"]
    for i, h in enumerate(headers):
        c = table_student.cell(0, i)
        set_cell_background(c, "2A2827")
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    student_fields = [
        ("id", "BigAuto (PK)", "Auto-increment", "Unique identifier for student"),
        ("email", "EmailField", "Unique, Required", "Student login email identifier"),
        ("phone_number", "CharField(20)", "Required", "Contact number (visible to confirmed roommates)"),
        ("first_name", "CharField(100)", "Required", "Student given name"),
        ("last_name", "CharField(100)", "Required", "Student surname / family name"),
        ("gender", "CharField(10)", "'MALE' | 'FEMALE'", "Student gender used for room policy enforcement"),
        ("program_of_study", "CharField(150)", "Optional / Blank", "Academic program (e.g. BSc Computer Science)"),
        ("year_of_study", "IntegerField", "1 to 6", "Academic level / year"),
        ("emergency_contact_name", "CharField(150)", "Optional / Blank", "Guardian or emergency contact name"),
        ("emergency_contact_phone", "CharField(20)", "Optional / Blank", "Guardian phone number"),
        ("is_active", "BooleanField", "Default: True", "Account active flag"),
        ("is_staff", "BooleanField", "Default: False", "Staff portal access flag"),
    ]
    for row in student_fields:
        r = table_student.add_row()
        for idx, val in enumerate(row):
            c = r.cells[idx]
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            p = c.paragraphs[0]
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Model 2: Room & RoomSpace
    add_section_header("Model: Room & RoomSpace (Inventory)", level=2)
    add_body_p(
        "Rooms represent physical living units. Bed spaces represent individual bookable slots within a room."
    )

    table_room = doc.add_table(rows=1, cols=4)
    table_room.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = table_room.cell(0, i)
        set_cell_background(c, "2A2827")
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    room_fields = [
        ("Room.room_number", "CharField(20)", "Unique, Indexed", "Room code: 'R01', 'R02', '101', '201'"),
        ("Room.floor", "IntegerField", "0 (Ground), 1, 2, 3", "Floor level of the hostel building"),
        ("Room.capacity", "IntegerField", "1, 2, or 3", "Total bed spaces inside the room"),
        ("Room.room_type", "CharField(20)", "'STANDARD' | 'PREMIUM'", "Room pricing & luxury tier"),
        ("Room.price_per_year", "DecimalField(10,2)", "Default: GHS 6,000", "Annual bed space fee"),
        ("Room.gender_policy", "CharField(15)", "'UNASSIGNED'|'MALE'|'FEMALE'", "Locks to first booking gender"),
        ("Room.amenities", "JSONField", "Default: list", "List of amenities (e.g. WiFi, AC, Balcony)"),
        ("RoomSpace.room", "FK -> Room", "on_delete=CASCADE", "Parent room relationship"),
        ("RoomSpace.space_number", "IntegerField", "1, 2, or 3", "Bed space identifier (e.g. Bed 1, Bed 2)"),
        ("RoomSpace.occupant", "FK -> Student", "Nullable, on_delete=SET_NULL", "Confirmed paid student occupant"),
        ("RoomSpace.is_held", "BooleanField", "Default: False", "Temporary 15-minute checkout lock flag"),
        ("RoomSpace.held_by", "FK -> Student", "Nullable, on_delete=SET_NULL", "Student currently holding the space"),
        ("RoomSpace.held_at", "DateTimeField", "Nullable", "Timestamp when hold was initiated"),
    ]
    for row in room_fields:
        r = table_room.add_row()
        for idx, val in enumerate(row):
            c = r.cells[idx]
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            p = c.paragraphs[0]
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Model 3: Booking & Payment
    add_section_header("Model: Booking & Payment (Transactions)", level=2)
    add_body_p(
        "Bookings record reservation lifecycles. Payment records track Paystack transaction references."
    )

    table_booking = doc.add_table(rows=1, cols=4)
    table_booking.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = table_booking.cell(0, i)
        set_cell_background(c, "2A2827")
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    booking_fields = [
        ("Booking.booking_id", "UUIDField", "Unique, Default: uuid4", "Public reservation identifier"),
        ("Booking.student", "FK -> Student", "CASCADE", "Student making the reservation"),
        ("Booking.space", "FK -> RoomSpace", "CASCADE", "Target bed space"),
        ("Booking.status", "CharField(20)", "'HOLD' | 'PENDING' | 'CONFIRMED' | 'CANCELLED' | 'EXPIRED'", "Current state of the reservation"),
        ("Booking.expires_at", "DateTimeField", "Indexed", "Hold expiration (now + 15 mins)"),
        ("Payment.booking", "OneToOne -> Booking", "CASCADE", "Linked reservation"),
        ("Payment.paystack_reference", "CharField(100)", "Unique, Indexed", "Paystack checkout reference"),
        ("Payment.amount", "DecimalField(10,2)", "GHS Amount", "Payment amount charged"),
        ("Payment.status", "CharField(20)", "'PENDING' | 'SUCCESS' | 'FAILED'", "Payment verification status"),
        ("Payment.channel", "CharField(50)", "card, mobile_money", "Payment channel used by student"),
    ]
    for row in booking_fields:
        r = table_booking.add_row()
        for idx, val in enumerate(row):
            c = r.cells[idx]
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            p = c.paragraphs[0]
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ------------------ SECTION 3: CORE BUSINESS RULES ------------------
    add_section_header("3. Core Business Rules & Security Guardrails")

    add_section_header("Rule 1: First-Come, First-Served Gender Locking", level=2)
    add_bullet_p(
        "All rooms initialize with gender_policy = 'UNASSIGNED'.",
        "Initial State: "
    )
    add_bullet_p(
        "When the very first student selects a bed space in an unassigned room and creates a hold, "
        "the backend atomically locks the room's gender_policy to match the student's gender ('MALE' or 'FEMALE').",
        "Lock Trigger: "
    )
    add_bullet_p(
        "Subsequent space selections in that room validate the requesting student's gender against the room policy. "
        "If a female student attempts to book a male-locked room, the API immediately responds with 400 Bad Request.",
        "Strict Isolation: "
    )
    add_bullet_p(
        "If all active holds in a room are released or expire without a confirmed payment, "
        "the room policy resets back to 'UNASSIGNED'.",
        "Auto-Reset: "
    )

    add_section_header("Rule 2: Concurrency-Safe 15-Minute Space Hold", level=2)
    add_bullet_p(
        "When a student selects a bed space, POST /api/bookings/hold/ is called. "
        "The backend sets RoomSpace.is_held = True, held_by = student, and held_at = now.",
        "Atomic Hold: "
    )
    add_bullet_p(
        "If another student simultaneously clicks the same bed space, the backend returns 409 Conflict. "
        "The frontend displays: 'This bed space is currently held by another student. Please try another space.'",
        "Conflict Handling: "
    )
    add_bullet_p(
        "Holds are valid for exactly 15 minutes. An active booking record is created with expires_at = now + 15m. "
        "The frontend starts a live 15-minute countdown timer.",
        "Duration: "
    )
    add_bullet_p(
        "If the user cancels the modal or clicks 'Change Space', POST /api/bookings/release-hold/ is triggered. "
        "If the student closes their browser tab, any queries after 15 minutes automatically treat the space as free.",
        "Release: "
    )

    add_section_header("Rule 3: Post-Payment Roommate Privacy Protection", level=2)
    add_bullet_p(
        "Students' personal information (names, phone numbers) must NEVER be exposed publicly or to guests.",
        "Privacy by Default: "
    )
    add_bullet_p(
        "Endpoint GET /api/bookings/roommates/<room_id>/ validates that the requesting student is an ACTIVE OCCUPANT "
        "of that specific room with a CONFIRMED payment. Unauthenticated requests or non-occupants receive 403 Forbidden.",
        "Permission Gating: "
    )
    add_bullet_p(
        "Once a student successfully books, their StudentDashboard invokes this endpoint and renders roommate contact cards "
        "with one-click phone links (tel:+233...).",
        "Frontend Render: "
    )

    add_section_header("Rule 4: HttpOnly Cookie JWT Authentication", level=2)
    add_bullet_p(
        "JWT tokens (access_token, refresh_token) are transmitted via HttpOnly, SameSite cookies. "
        "JavaScript cannot read these tokens, preventing XSS token theft.",
        "Cookie Security: "
    )
    add_bullet_p(
        "Axios is configured with withCredentials: true on every request, transmitting cookies automatically.",
        "Client Setup: "
    )

    # ------------------ SECTION 4: API ENDPOINT SPECIFICATION ------------------
    add_section_header("4. Complete API Contract for Frontend")
    add_body_p(
        "Base URL in local development: http://localhost:8000/api\n"
        "Base URL behind Nginx / Production: /api"
    )

    table_api = doc.add_table(rows=1, cols=4)
    table_api.alignment = WD_TABLE_ALIGNMENT.CENTER
    api_headers = ["Method & Endpoint", "Auth", "Payload / Query Params", "Response"]
    for i, h in enumerate(api_headers):
        c = table_api.cell(0, i)
        set_cell_background(c, "2A2827")
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    api_endpoints = [
        ("POST /api/auth/register/", "Public", "{\n  email, password,\n  first_name, last_name,\n  phone_number, gender\n}", "201 Created\nSets HttpOnly cookies\nReturns student profile object"),
        ("POST /api/auth/login/", "Public", "{\n  email, password\n}", "200 OK\nSets HttpOnly cookies\nReturns student profile object"),
        ("GET /api/auth/me/", "Cookie", "None", "200 OK: Current logged-in student\n401 Unauthorized: Guest"),
        ("POST /api/auth/logout/", "Cookie", "None", "200 OK: Clears JWT cookies"),
        ("GET /api/rooms/", "Public", "?floor=1&capacity=2\n&gender=MALE&available_only=true", "200 OK: Array of 102 room objects with embedded spaces list & occupancy count"),
        ("GET /api/rooms/<id>/", "Public", "None", "200 OK: Single room detail with spaces"),
        ("POST /api/bookings/hold/", "Required", "{\n  space_id: number,\n  room_number: string,\n  space_number: number\n}", "201 Created: { booking_id, expires_at }\n409 Conflict: Space already held\n400: Gender mismatch"),
        ("POST /api/bookings/release-hold/", "Required", "{\n  booking_id: string\n}", "200 OK: { detail: 'Hold released' }"),
        ("POST /api/bookings/initialize-payment/", "Required", "{\n  booking_id: string\n}", "200 OK: {\n  reference: string,\n  authorization_url: string,\n  amount: number\n}"),
        ("POST /api/bookings/verify-payment/", "Required", "{\n  reference: string\n}", "200 OK: Booking confirmed, space assigned, room gender locked"),
        ("GET /api/bookings/roommates/<room_id>/", "Required", "room_id: string | number", "200 OK: Array of confirmed roommates with name and phone.\n403 Forbidden: Not occupant"),
        ("POST /api/bookings/webhook/", "Paystack", "Paystack charge.success IPN", "200 OK (Verifies HMAC-SHA512 header)"),
    ]
    for row in api_endpoints:
        r = table_api.add_row()
        for idx, val in enumerate(row):
            c = r.cells[idx]
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            p = c.paragraphs[0]
            run = p.add_run(val)
            run.font.name = "Consolas" if idx in [0, 2] else "Arial"
            run.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ------------------ SECTION 5: STEP-BY-STEP WORKFLOWS ------------------
    add_section_header("5. End-to-End User Workflows (Walkflows)")

    add_section_header("Workflow A: User Registration & Authentication", level=2)
    add_bullet_p("User clicks 'Sign In' or 'Join Portal' in the navigation bar.", "Step 1: ")
    add_bullet_p("LoginRegisterModal opens. User enters email, password, full name, phone number, and selects Gender (Male/Female).", "Step 2: ")
    add_bullet_p("Frontend calls registerUser(payload) in src/services/api.ts.", "Step 3: ")
    add_bullet_p("Backend creates Student record, generates JWT tokens, sets HttpOnly cookies, and returns student profile.", "Step 4: ")
    add_bullet_p("HostelContext updates currentUser state; Navbar switches to show student name and 'My Dashboard' button.", "Step 5: ")

    add_section_header("Workflow B: Room Browsing & Filter Discovery", level=2)
    add_bullet_p("User navigates to the Room Catalog section.", "Step 1: ")
    add_bullet_p("HostelContext fetches GET /api/rooms/. All 102 rooms and 193 bed spaces are loaded into memory.", "Step 2: ")
    add_bullet_p("User filters by Floor (Ground, 1, 2, 3), Room Type (Standard, Premium), or Gender Policy.", "Step 3: ")
    add_bullet_p("Each card displays live occupancy status (e.g. '1/2 Beds Taken', 'Male Room', 'Unassigned').", "Step 4: ")

    add_section_header("Workflow C: Space Selection & 15-Minute Hold", level=2)
    add_bullet_p("User clicks 'Select Room' to open RoomDetailModal.", "Step 1: ")
    add_bullet_p("User clicks on an available bed space (e.g. Bed Space #2).", "Step 2: ")
    add_bullet_p("If the user is not logged in, LoginRegisterModal is prompted automatically.", "Step 3: ")
    add_bullet_p("HostelContext calls lockSpaceTemporarily(spaceId). Backend executes atomic transaction: checks gender compatibility, verifies space is not occupied or held, creates a Booking with status 'HOLD', and starts 15-minute timer.", "Step 4: ")
    add_bullet_p("If another student has held the space, 409 is returned and an error toast is displayed. Otherwise, BookingFlowModal opens displaying the 15:00 countdown clock.", "Step 5: ")

    add_section_header("Workflow D: Checkout & Paystack Payment", level=2)
    add_bullet_p("User reviews booking summary (Room Number, Space Number, GHS 6,000 fee).", "Step 1: ")
    add_bullet_p("User clicks 'Proceed to Pay with Paystack'.", "Step 2: ")
    add_bullet_p("Frontend calls initializePayment(bookingId). Backend returns Paystack reference and authorization_url.", "Step 3: ")
    add_bullet_p("User completes payment via Paystack popup / modal using Mobile Money (MTN, Telecel, AT) or Visa/Mastercard.", "Step 4: ")
    add_bullet_p("On success callback, frontend calls verifyPayment(reference).", "Step 5: ")
    add_bullet_p("Backend verifies transaction with Paystack API, sets Booking to 'CONFIRMED', permanently assigns RoomSpace.occupant = student, and locks Room.gender_policy. Returns success receipt.", "Step 6: ")
    add_bullet_p("ReceiptModal opens showing downloadable/printable payment confirmation with transaction reference.", "Step 7: ")

    add_section_header("Workflow E: Student Dashboard & Roommate Discovery", level=2)
    add_bullet_p("Student opens StudentDashboard from the navigation bar.", "Step 1: ")
    add_bullet_p("Dashboard displays the confirmed reservation: Room Number, Assigned Bed Space, Floor, and Academic Year.", "Step 2: ")
    add_bullet_p("Dashboard triggers getRoommates(roomId). Since the user is the confirmed occupant, backend responds with 200 OK and roommate details.", "Step 3: ")
    add_bullet_p("Dashboard renders Roommate Cards displaying their full name and a direct clickable phone number (tel:024XXXXXXX) for WhatsApp or calling.", "Step 4: ")

    # ------------------ SECTION 6: FRONTEND ARCHITECTURE & CONTEXT ------------------
    add_section_header("6. Frontend State Management & Code Structure")
    add_body_p(
        "All state logic is centralized in src/context/HostelContext.tsx. "
        "Here is the map of frontend components and their backend connections:"
    )

    table_comp = doc.add_table(rows=1, cols=3)
    table_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    comp_headers = ["Component File", "State / Hook", "Backend API Integration"]
    for i, h in enumerate(comp_headers):
        c = table_comp.cell(0, i)
        set_cell_background(c, "2A2827")
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    components_list = [
        ("src/services/api.ts", "Axios Instance", "Dynamic baseURL fallback, withCredentials: true, all API helper methods"),
        ("src/context/HostelContext.tsx", "useHostel() Context", "Global rooms cache, activeHold, currentUser, booking flow orchestration"),
        ("src/components/LoginRegisterModal.tsx", "isAuthModalOpen", "Calls registerUser() and loginUser(), updates currentUser"),
        ("src/components/Navbar.tsx", "currentUser, logout", "Displays user profile pill, triggers login modal or logoutUser()"),
        ("src/components/RoomCatalog.tsx", "rooms, filters", "Consumes backend rooms list, renders RoomCard components"),
        ("src/components/RoomDetailModal.tsx", "selectedRoom, lockSpace", "Displays bed spaces layout, triggers lockSpaceTemporarily()"),
        ("src/components/BookingFlowModal.tsx", "activeHold, timer", "15-minute countdown, Paystack initialization & verification"),
        ("src/components/ReceiptModal.tsx", "latestBooking", "Displays confirmed booking receipt and payment reference"),
        ("src/components/StudentDashboard.tsx", "userBooking, getRoommates", "Fetches private roommate contacts and displays active room"),
        ("src/components/AdminDashboard.tsx", "Direct Admin Link", "One-click link to Django Admin (/admin/) for full staff controls"),
    ]
    for row in components_list:
        r = table_comp.add_row()
        for idx, val in enumerate(row):
            c = r.cells[idx]
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            p = c.paragraphs[0]
            run = p.add_run(val)
            run.font.name = "Consolas" if idx == 0 else "Arial"
            run.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ------------------ SECTION 7: NGINX & DEPLOYMENT ------------------
    add_section_header("7. Nginx Reverse Proxy & Production Deployment")
    add_body_p(
        "A complete Nginx configuration file is ready at mush-backend/nginx.conf. "
        "Here is why it is used and how it routes traffic:"
    )

    add_code_block("""# Nginx Unified Origin Architecture
# mush-backend/nginx.conf

upstream django_backend {
    server 127.0.0.1:8000;
}

upstream vite_frontend {
    server 127.0.0.1:3000;
}

server {
    listen 80;
    server_name localhost mushiahostel.com www.mushiahostel.com;

    # 1. Django REST API
    location /api/ {
        proxy_pass http://django_backend;
        proxy_set_header Host $host;
        proxy_set_header X-Real-IP $remote_addr;
        proxy_pass_header Set-Cookie;
    }

    # 2. Django Admin Portal
    location /admin/ {
        proxy_pass http://django_backend;
        proxy_set_header Host $host;
    }

    # 3. Static & Media Files (Direct Disk Delivery with 30-day Cache)
    location /static/ {
        alias /var/www/mush-backend/static/;
        expires 30d;
    }
    location /media/ {
        alias /var/www/mush-backend/media/;
        expires 30d;
    }

    # 4. React Vite Frontend
    location / {
        proxy_pass http://vite_frontend;
        proxy_set_header Upgrade $http_upgrade;
        proxy_set_header Connection "upgrade";
    }
}""")

    add_section_header("Production Deployment Quick-Start Commands", level=2)
    add_bullet_p("cd /path/to/MUSH && pnpm install && pnpm build", "1. Build Frontend: ")
    add_bullet_p("cd /path/to/mush-backend && uv run python manage.py collectstatic --noinput", "2. Collect Static: ")
    add_bullet_p("uv run gunicorn core.wsgi:application --bind 127.0.0.1:8000 --workers 3", "3. Start Django Server: ")
    add_bullet_p("sudo cp nginx.conf /etc/nginx/sites-available/mushia.conf && sudo systemctl restart nginx", "4. Start Nginx: ")

    add_section_header("Default Credentials Reference", level=2)
    add_bullet_p("admin", "Admin Username: ")
    add_bullet_p("admin123", "Admin Password: ")
    add_bullet_p("http://localhost:8000/admin/", "Admin Portal URL: ")

    # Save to both paths
    output_path_desktop = r"C:\Users\user\Desktop\Mushia_Hostel_Frontend_Integration_Guide.docx"
    output_path_frontend = r"C:\Users\user\Desktop\MUSH\Mushia_Hostel_Frontend_Integration_Guide.docx"
    output_path_backend = r"C:\Users\user\Desktop\mush-backend\Mushia_Hostel_Frontend_Integration_Guide.docx"

    doc.save(output_path_desktop)
    doc.save(output_path_frontend)
    doc.save(output_path_backend)

    print(f"Documentation saved successfully to:\n- {output_path_desktop}\n- {output_path_frontend}\n- {output_path_backend}")

if __name__ == "__main__":
    create_documentation()
